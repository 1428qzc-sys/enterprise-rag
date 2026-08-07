"""多格式文档解析：PDF / Word / Excel / Markdown / TXT / 网页 URL。

统一产出 ``Section`` 列表（正文 + 可选页码 / 小标题），页码信息会一路带到
chunk 元数据里，用于问答时的精确溯源（例如「PDF 第 3 页」）。
"""

from __future__ import annotations

import os
import zipfile
from dataclasses import dataclass, field
from typing import List, Optional

from ..config import settings
from ..security_network import fetch_public_url


@dataclass
class Section:
    text: str
    page: Optional[int] = None
    title: Optional[str] = None
    extra: dict = field(default_factory=dict)


SUPPORTED_EXTS = {".pdf", ".docx", ".xlsx", ".md", ".markdown", ".txt", ".csv", ".htm", ".html"}

CANONICAL_MIME = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".txt": "text/plain",
    ".csv": "text/csv",
    ".htm": "text/html",
    ".html": "text/html",
}


def validate_file_content(path: str, filename: str) -> str:
    """校验扩展名与文件结构，返回可信的规范 MIME。"""

    ext = os.path.splitext(filename.lower())[1]
    if ext not in SUPPORTED_EXTS:
        raise ValueError(f"不支持的文件类型：{ext or '未知'}")
    with open(path, "rb") as source:
        head = source.read(8192)
    if not head:
        raise ValueError("文件为空")
    if ext == ".pdf" and not head.startswith(b"%PDF-"):
        raise ValueError("PDF 文件头无效")
    if ext in {".docx", ".xlsx"}:
        if not zipfile.is_zipfile(path):
            raise ValueError(f"{ext[1:].upper()} 文件结构无效")
        required = "word/document.xml" if ext == ".docx" else "xl/workbook.xml"
        with zipfile.ZipFile(path) as archive:
            if required not in archive.namelist():
                raise ValueError(f"{ext[1:].upper()} 缺少必需内容")
    if ext in {".md", ".markdown", ".txt", ".csv", ".htm", ".html"}:
        if b"\x00" in head:
            raise ValueError("文本文件包含二进制内容")
        if not any(
            _can_decode(head, encoding)
            for encoding in ("utf-8", "utf-8-sig", "gbk", "latin-1")
        ):
            raise ValueError("文本文件编码无法识别")
    return CANONICAL_MIME[ext]


def _can_decode(raw: bytes, encoding: str) -> bool:
    try:
        raw.decode(encoding)
        return True
    except UnicodeDecodeError:
        return False


def detect_type(filename: str) -> str:
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext == ".xlsx":
        return "xlsx"
    if ext in (".htm", ".html"):
        return "html"
    if ext in (".md", ".markdown"):
        return "markdown"
    return "text"  # txt / csv / 其它按纯文本


def _read_text_smart(path: str) -> str:
    """按 utf-8 → gbk → latin-1 尝试解码，兼容中文 Windows 文本。"""
    raw = open(path, "rb").read()
    for enc in ("utf-8", "utf-8-sig", "gbk", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="ignore")


# ---------------- 各格式解析 ----------------
def _parse_pdf(path: str) -> List[Section]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    sections: List[Section] = []
    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()
        if text:
            sections.append(Section(text=text, page=i + 1))
    return sections


def _parse_docx(path: str) -> List[Section]:
    import docx

    doc = docx.Document(path)
    parts: List[str] = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    return [Section(text=text)] if text else []


def _parse_xlsx(path: str) -> List[Section]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sections: List[Section] = []
    for ws in wb.worksheets:
        lines: List[str] = []
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(c.strip() for c in cells):
                lines.append(" | ".join(cells))
        text = "\n".join(lines).strip()
        if text:
            sections.append(Section(text=text, title=ws.title))
    wb.close()
    return sections


def _parse_html_text(html: str) -> tuple[str, Optional[str]]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "noscript", "header", "footer", "nav", "aside"]):
        tag.decompose()
    title = soup.title.string.strip() if soup.title and soup.title.string else None
    text = soup.get_text(separator="\n")
    lines = [ln.strip() for ln in text.splitlines()]
    text = "\n".join(ln for ln in lines if ln)
    return text, title


def _parse_html_file(path: str) -> List[Section]:
    text, title = _parse_html_text(_read_text_smart(path))
    return [Section(text=text, title=title)] if text.strip() else []


def _parse_plain(path: str) -> List[Section]:
    text = _read_text_smart(path).strip()
    return [Section(text=text)] if text else []


def parse_file(path: str, filename: Optional[str] = None) -> List[Section]:
    """按文件类型解析为 Section 列表。"""
    filename = filename or os.path.basename(path)
    kind = detect_type(filename)
    if kind == "pdf":
        return _parse_pdf(path)
    if kind == "docx":
        return _parse_docx(path)
    if kind == "xlsx":
        return _parse_xlsx(path)
    if kind == "html":
        return _parse_html_file(path)
    return _parse_plain(path)


def parse_url(url: str, timeout: float = 30.0) -> tuple[List[Section], str]:
    """抓取网页并抽取正文。返回 (sections, 页面标题)。"""
    del timeout  # 使用集中配置，避免调用方绕过安全基线。
    fetched = fetch_public_url(
        url,
        timeout=settings.url_fetch_timeout_seconds,
        max_bytes=settings.url_fetch_max_mb * 1024 * 1024,
        max_redirects=settings.url_fetch_max_redirects,
    )
    html = fetched.text
    text, title = _parse_html_text(html)
    title = title or url
    sections = [Section(text=text, title=title)] if text.strip() else []
    return sections, title
